# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

INK=RGBColor(0x29,0x25,0x24); GREEN=RGBColor(0x16,0x65,0x34); MUST=RGBColor(0xD9,0x77,0x06)
SLATE=RGBColor(0x47,0x53,0x69); MUT=RGBColor(0x78,0x71,0x6C); LINE=RGBColor(0xE7,0xE5,0xE4)
WHITE=RGBColor(0xFF,0xFF,0xFF); LT=RGBColor(0xD6,0xD3,0xD1); INKBODY=RGBColor(0x44,0x40,0x3C)
HEAD="Cambria"; BODY="Calibri"

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BL=prs.slide_layouts[6]

def slide(bg=WHITE):
    s=prs.slides.add_slide(BL)
    r=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb=bg; r.line.fill.background(); r.shadow.inherit=False
    return s

def box(s,x,y,w,h,text,size,color,bold=False,font=BODY,align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP,italic=False,sp=None):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame
    tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    for i,ln in enumerate(text.split("\n")):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align
        if sp: p.line_spacing=sp
        r=p.add_run(); r.text=ln; f=r.font
        f.size=Pt(size); f.bold=bold; f.italic=italic; f.name=font; f.color.rgb=color
    return tb

def rrect(s,x,y,w,h,fill,radius=0.08,lcolor=None):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=fill; sh.shadow.inherit=False
    if lcolor: sh.line.color.rgb=lcolor; sh.line.width=Pt(1)
    else: sh.line.fill.background()
    try: sh.adjustments[0]=radius
    except Exception: pass
    return sh

def oval(s,x,y,d,fill):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(d),Inches(d))
    sh.fill.solid(); sh.fill.fore_color.rgb=fill; sh.line.fill.background(); sh.shadow.inherit=False
    return sh

def badge(s,x,y,num,color=GREEN):
    oval(s,x,y,0.62,color)
    box(s,x,y-0.02,0.62,0.66,str(num),22,WHITE,bold=True,font=HEAD,
        align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)

def mark(s,x,y,scale=1.0):
    t=0.9*scale
    rrect(s,x,y,t,t,GREEN,radius=0.26)
    box(s,x,y-0.04*scale,t,t,"P",int(46*scale),WHITE,bold=True,font=HEAD,
        align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    oval(s,x+t-0.20*scale,y+0.06*scale,0.16*scale,MUST)

def stat(s,x,y,w,big,label,color=GREEN,bigsize=46):
    box(s,x,y,w,1.0,big,bigsize,color,bold=True,font=HEAD,align=PP_ALIGN.LEFT)
    box(s,x,y+0.92,w,0.9,label,13.5,MUT,font=BODY,align=PP_ALIGN.LEFT,sp=1.05)

def header(s,num,title):
    badge(s,0.7,0.62,num)
    box(s,1.5,0.6,11.1,0.95,title,33,INK,bold=True,font=HEAD,anchor=MSO_ANCHOR.MIDDLE)

ND="–"  # en dash
RA="→"  # arrow
AP="’"  # apostrophe
AX="≈"  # approx

# 1 TITLE
s=slide(INK)
mark(s,0.95,0.85,1.0)
box(s,1.95,0.92,6,0.8,"PRIMUS",26,WHITE,bold=True,font=HEAD,anchor=MSO_ANCHOR.MIDDLE)
box(s,1.95,1.45,8,0.5,"PERSON-YEAR FORECASTS",12,LT,font=BODY,anchor=MSO_ANCHOR.MIDDLE)
box(s,0.95,2.9,11.4,1.4,"Why believe it",60,WHITE,bold=True,font=HEAD)
box(s,1.0,4.25,11,0.7,"Speaker 2 "+ND+" the method behind the numbers",22,MUST,font=BODY)
box(s,1.0,5.5,11,0.6,"Stops 7"+ND+"10  ·  ~9 minutes  ·  Longevity & the Labour Market",15,LT,font=BODY)
box(s,1.0,6.5,11,0.5,"Part 2 of 2 "+ND+" after the findings, here is why you can trust them.",13.5,LT,font=BODY,italic=True)

# 2 METHODOLOGY
s=slide()
header(s,7,"Methodology "+ND+" how it works")
box(s,1.5,1.55,7,0.6,"Five simple models, averaged "+ND+" deliberately no neural networks.",17,INKBODY,font=BODY)
cy=2.45
for m in ["Linear trend","Damped Holt","Drift","Na"+chr(0xEF)+"ve","AR(1) "+ND+" mean-reverting"]:
    rrect(s,1.5,cy,4.4,0.62,RGBColor(0xF0,0xF5,0xF1),radius=0.3)
    oval(s,1.72,cy+0.19,0.24,GREEN)
    box(s,2.15,cy,3.6,0.62,m,15,INK,bold=True,font=BODY,anchor=MSO_ANCHOR.MIDDLE)
    cy+=0.78
rrect(s,6.6,2.45,6.1,3.9,RGBColor(0xFB,0xFB,0xFA),radius=0.05,lcolor=LINE)
stat(s,7.0,2.85,5.4,"3","training pairs a dedicated 9-year model gets "+ND+" so it starves and overfits.",MUST,bigsize=52)
stat(s,7.0,4.55,5.4,"12"+ND+"40%","honest band width: model spread + real backtest error + migration.",GREEN,bigsize=40)
box(s,1.5,6.55,11.2,0.7,"We tested it: per-horizon specialists lose at every horizon "+ND+" and the na"+chr(0xEF)+"ve last-value is the hardest thing to beat, so we win by restraint.",13.5,MUT,font=BODY,italic=True,sp=1.05)

# 3 PLAYGROUND
s=slide()
header(s,8,"Model playground "+ND+" drag it live")
box(s,1.5,1.55,11,0.6,"Two interactive toys make the method tangible.",17,INKBODY,font=BODY)
rrect(s,1.5,2.4,5.3,3.5,RGBColor(0xFB,0xFB,0xFA),radius=0.05,lcolor=LINE)
box(s,1.85,2.7,4.7,0.6,"Ensemble uncertainty",18,GREEN,bold=True,font=HEAD)
box(s,1.85,3.4,4.7,2.3,"Add more simple models and a little noise "+ND+" a cloud of forecasts forms, and the band is just its spread.\n\nUncertainty = the disagreement among honest simple models.",14.5,INKBODY,font=BODY,sp=1.1)
rrect(s,7.05,2.4,5.65,3.5,RGBColor(0xFB,0xFB,0xFA),radius=0.05,lcolor=LINE)
box(s,7.4,2.7,5,0.6,"Flexibility "+ND+" bias vs variance",18,GREEN,bold=True,font=HEAD)
box(s,7.4,3.35,5,0.55,"Cross-validated error as you over-flex:",13.5,MUT,font=BODY)
box(s,7.4,3.85,5.2,1.0,"0.7  "+RA+"  15.6",44,MUST,bold=True,font=HEAD)
box(s,7.4,4.9,5,0.9,"Training error keeps falling, but CV error explodes "+ND+" overfitting, live. Sweet spot: degree 3.",14.5,INKBODY,font=BODY,sp=1.05)
box(s,1.5,6.35,11.2,0.8,"That is exactly why we use simple models "+ND+" a neural network would land far right, deep in the overfit zone.",13.5,MUT,font=BODY,italic=True,sp=1.05)

# 4 TRUST FOR
s=slide()
header(s,9,"Trust "+ND+" the case for it")
box(s,1.5,1.55,11,0.6,"Rolling-origin cross-validation: hold out the last years, forecast from history only.",17,INKBODY,font=BODY)
for x,(v,l,c) in zip([1.5,5.25,9.0],[("3.6%","Supply forecast error (MAPE)",GREEN),("2.2%","Demand forecast error (MAPE)",GREEN),("+0.29","CRPS skill vs na"+chr(0xEF)+"ve "+ND+" every driver",MUST)]):
    rrect(s,x,2.6,3.4,2.2,RGBColor(0xFB,0xFB,0xFA),radius=0.06,lcolor=LINE)
    box(s,x+0.35,2.9,2.7,1.0,v,52,c,bold=True,font=HEAD)
    box(s,x+0.35,3.95,2.7,0.8,l,13.5,MUT,font=BODY,sp=1.05)
box(s,1.5,5.5,11.2,1.3,"A proper scoring rule (CRPS) grades the whole band, not just the point "+ND+" and it is positive for every driver, even vacancies whose point forecast barely ties na"+chr(0xEF)+"ve. The bands are well-calibrated: the model knows what it doesn"+AP+"t know.",15,INKBODY,font=BODY,sp=1.12)

# 5 CRITIQUE
s=slide()
header(s,9,"…and the case against "+ND+" our own audit")
cards=[("Honest precision","Bootstrap confidence intervals on every number.","Vacancies: 24%  "+RA+"  [16"+ND+"34]"),
       ("The fragile headline","Strip out non-comparable self-reported health and the East"+ND+"West split nearly collapses.","West  "+ND+"287  "+RA+"  +642 M"),
       ("0.99 is fake, "+AX+"0 is real","Cross-country R"+chr(0xB2)+" vs within-country R"+chr(0xB2)+" for longevity "+RA+" spending.","0.99   vs   "+AX+" 0.00")]
x=1.5
for t,d,k in cards:
    rrect(s,x,2.0,3.6,3.7,WHITE,radius=0.05,lcolor=LINE)
    box(s,x+0.3,2.25,3.05,0.85,t,16.5,SLATE,bold=True,font=HEAD,sp=0.98)
    box(s,x+0.3,3.2,3.0,1.55,d,13.5,INKBODY,font=BODY,sp=1.08)
    box(s,x+0.3,4.9,3.05,0.7,k,16,MUST,bold=True,font=HEAD)
    x+=3.83
box(s,1.5,6.15,11.2,0.8,"And every fancier tool we measured "+ND+" a regression tree, gradient boosting, per-horizon specialists "+ND+" lost. We didn"+AP+"t just prefer simple; we proved simple wins.",13.5,MUT,font=BODY,italic=True,sp=1.05)

# 6 WHAT-IF
s=slide()
header(s,10,"What-if sandbox "+ND+" the finale")
box(s,1.5,1.55,11,0.6,"A live engine: three dials, four answers that update instantly.",17,INKBODY,font=BODY)
dials=[("Healthy life years","stay healthy longer "+RA+" supply rises"),("Retirement age","later "+RA+" who it helps, who it strains"),("Participation","more workers "+RA+" the shortage shrinks")]
x=1.5
for i,(t,d) in enumerate(dials):
    rrect(s,x,2.5,3.6,2.3,RGBColor(0xF0,0xF5,0xF1),radius=0.06)
    oval(s,x+0.3,2.8,0.5,GREEN); box(s,x+0.3,2.78,0.5,0.5,str(i+1),16,WHITE,bold=True,font=HEAD,align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    box(s,x+0.95,2.78,2.5,0.55,t,15.5,INK,bold=True,font=HEAD,anchor=MSO_ANCHOR.MIDDLE)
    box(s,x+0.3,3.55,3.05,1.1,d,14,INKBODY,font=BODY,sp=1.1)
    x+=3.83
box(s,1.5,5.4,11.2,1.4,"Who keeps working, who retires, who needs care, and who fills the concert halls and the adult-education classes. Not a prophecy "+ND+" a transparent engine you can interrogate, reproducible from raw data to this slider.",15,INKBODY,font=BODY,sp=1.12)

# 7 CLOSING
s=slide(INK)
box(s,1.0,1.7,11.4,2.2,"The longevity dividend is real "+ND+" but it sits on a deep East"+ND+"West divide.",40,WHITE,bold=True,font=HEAD,sp=1.05)
box(s,1.0,4.0,11.0,1.2,"Honest about what it knows and what it doesn"+AP+"t. Reproducible end to end. Trust the direction "+ND+" discount the decimals.",18,LT,font=BODY,sp=1.15)
mark(s,1.0,5.9,0.8)
box(s,1.78,5.95,7,0.7,"Thank you "+ND+" questions?",22,MUST,bold=True,font=HEAD,anchor=MSO_ANCHOR.MIDDLE)

import os
out="C:/Users/Dell/Desktop/SUMMER/docs/Primus_Speaker2_WhyBelieveIt.pptx"
os.makedirs(os.path.dirname(out),exist_ok=True)
prs.save(out)
print("saved", out, "| slides:", len(prs.slides._sldIdLst))
